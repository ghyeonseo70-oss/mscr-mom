"""8월 4째주 랩미팅 pptx 조립 - s/phi/L_M/Fx/Fy 5개 타겟 성능 진단 결과와, 그중 안 풀린
문제(Fx_board 고각도, s 팁근처, L_M=0)에 대해 시도했다가 실패/부분성공한 내용을 정리.
스타일은 기존 '랩미팅_8월1째주 (2).pptx'/build_lab_meeting_0821.py에서 쓰던 값 그대로 사용.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

GREEN = RGBColor(0x27, 0xAE, 0x60)
TITLE_DARK = RGBColor(0x22, 0x22, 0x22)
NAVY = RGBColor(0x1F, 0x2A, 0x44)
GRAY = RGBColor(0x6B, 0x72, 0x80)
BOX_FILL = RGBColor(0xEA, 0xF3, 0xEC)
RED = RGBColor(0xC0, 0x39, 0x2B)

EMU_PER_IN = 914400
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

HERE = os.path.dirname(__file__)
FEA_DIR = os.path.join(HERE, "fea")
CS_DIR = os.path.join(HERE, "..", "..", "data", "contact_scenarios")
FORCE_MODEL_DIR = os.path.join(HERE, "..", "..", "..", "data", "force_model")
OUT_PATH = r"C:\Users\gustj\Desktop\김현서\랩미팅 발표 자료\랩미팅_8월4째주.pptx"


def in_(v):
    return Emu(int(v * EMU_PER_IN))


def set_run(run, size_pt, color, bold=False, font=None, align=None):
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    if font:
        run.font.name = font


def add_bottom_line(slide):
    line = slide.shapes.add_connector(1, in_(0.4), in_(7.1484), in_(12.9), in_(7.1484))
    line.line.color.rgb = GREEN
    line.line.width = Emu(19050)


def add_page_number(slide, n):
    tb = slide.shapes.add_textbox(in_(11.85), in_(7.05), in_(0.9), in_(0.35))
    tf = tb.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    run = tf.paragraphs[0].add_run()
    run.text = str(n)
    set_run(run, 11, GRAY)


def add_title_bar(slide, text):
    rect = slide.shapes.add_shape(1, in_(0.45), in_(0.35), in_(0.09), in_(0.55))
    rect.fill.solid()
    rect.fill.fore_color.rgb = GREEN
    rect.line.fill.background()

    tb = slide.shapes.add_textbox(in_(0.65), in_(0.28), in_(11.5), in_(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_run(run, 26, TITLE_DARK, bold=True)


def new_content_slide(prs, title, page_num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, title)
    add_bottom_line(slide)
    add_page_number(slide, page_num)
    return slide


def add_picture_fit(slide, path, max_w_in, max_h_in, top_in, center_x_in=6.667):
    from PIL import Image
    w_px, h_px = Image.open(path).size
    aspect = w_px / h_px
    w_in, h_in = max_w_in, max_w_in / aspect
    if h_in > max_h_in:
        h_in = max_h_in
        w_in = max_h_in * aspect
    left = center_x_in - w_in / 2
    slide.shapes.add_picture(path, in_(left), in_(top_in), width=in_(w_in), height=in_(h_in))


def add_box(slide, left_in, top_in, w_in, h_in, text, fill=BOX_FILL, text_color=TITLE_DARK,
            size_pt=15, bold=False, line_color=GREEN):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, in_(left_in), in_(top_in), in_(w_in), in_(h_in))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = line_color
    box.line.width = Emu(19050)
    tf = box.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line
        set_run(run, size_pt, text_color, bold=bold)
    return box


def add_bullets(slide, left_in, top_in, w_in, h_in, lines, size_pt=17):
    tb = slide.shapes.add_textbox(in_(left_in), in_(top_in), in_(w_in), in_(h_in))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (line, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = "• " + line
        set_run(run, size_pt, RED if bold else TITLE_DARK, bold=bold)
        p.space_after = Pt(10)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ---- Title slide ----
    slide = prs.slides.add_slide(blank)
    rect = slide.shapes.add_shape(1, in_(0.5), in_(0.45), in_(0.35), in_(0.7))
    rect.fill.solid()
    rect.fill.fore_color.rgb = GREEN
    rect.line.fill.background()

    tb = slide.shapes.add_textbox(in_(0.5), in_(2.5), in_(11), in_(0.85))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "8월 4째주 랩미팅"
    set_run(run, 44, GREEN, bold=True)

    tb = slide.shapes.add_textbox(in_(0.5), in_(3.4), in_(11), in_(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = "5개 예측 타겟(s, phi, L_M, Fx, Fy) 성능 진단 — 뭐가 왜 안 되는지, 그리고 고쳐보려 한 시도들"
    set_run(run, 18, TITLE_DARK)

    tb = slide.shapes.add_textbox(in_(0.5), in_(4.7), in_(10), in_(0.5))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "20242260 김현서"
    set_run(run, 20, NAVY, bold=True)

    tb = slide.shapes.add_textbox(in_(0.5), in_(5.2), in_(10), in_(0.5))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "전기전자공학전공"
    set_run(run, 14, GRAY)

    add_bottom_line(slide)

    # ---- Slide 2: 전체 요약 ----
    slide = new_content_slide(prs, "1. 전체 요약 — 5개 타겟 중 어디가 문제인가", 2)
    add_picture_fit(slide, os.path.join(CS_DIR, "final_scatter_summary_0827.png"),
                     max_w_in=11.8, max_h_in=4.5, top_in=1.25)
    add_bullets(slide, 0.6, 5.75, 11.2, 1.3, [
        ("phi(R²=0.946), Fy_board(R²=0.877), s(R²=0.871)는 양호 — 실측 홀드아웃(n=99) 기준", False),
        ("L_M은 원래 안 됐지만(순수회귀 R²=0.718) 분류+회귀 하이브리드 구조로 R²=0.809까지 개선", False),
        ("Fx_board만 아직 약함(R²=0.600, n=99) — 다음 슬라이드부터 원인/시도 정리", True),
    ], size_pt=15)

    # ---- Slide 3: Fx_board 원인 ----
    slide = new_content_slide(prs, "2. Fx_board 문제 원인 — 왜 |phi|>=90에서 무너지는가", 3)
    add_picture_fit(slide, os.path.join(FORCE_MODEL_DIR, "fx_fy_geometry_example_phi30.png"),
                     max_w_in=6.6, max_h_in=4.9, top_in=1.35, center_x_in=3.55)
    add_picture_fit(slide, os.path.join(CS_DIR, "phi_breakdown_fx_bar.png"),
                     max_w_in=5.0, max_h_in=4.9, top_in=1.35, center_x_in=9.6)
    add_bullets(slide, 0.6, 6.35, 11.6, 0.8, [
        ("Fx_board는 물리적으로 F·sinθ 성분 — θ가 0을 넘나들 때마다 부호가 뒤집혀 원래도 어려운 "
         "타겟인데, |phi|>=90에서 특히 심함(R²=0.782 -> 0.355)", False),
    ], size_pt=14)

    # ---- Slide 4: Fx_board 해결 시도 (실패) ----
    slide = new_content_slide(prs, "3. Fx_board 해결 시도 — HIGH_PHI_WEIGHT 재튜닝, 해결 안 됨", 4)
    add_picture_fit(slide, os.path.join(CS_DIR, "high_phi_weight_attempts_bar.png"),
                     max_w_in=7.5, max_h_in=4.9, top_in=1.3, center_x_in=6.667)
    add_bullets(slide, 0.6, 6.35, 11.6, 0.8, [
        ("|phi|>=90 구간 힘 손실 가중치를 3.0으로 줘봤지만 그 구간 R²=0.343으로 그대로 -> 1.5로 "
         "낮춰 재시도해도 0.355로 사실상 무변화 -> 이 레버는 폐기, 실측 FEA 추가로 방향 전환", True),
    ], size_pt=14)

    # ---- Slide 5: s 문제 ----
    slide = new_content_slide(prs, "4. s 문제 — 팁 근처(60-100mm)에서 오차 증가", 5)
    add_picture_fit(slide, os.path.join(CS_DIR, "s_error_by_bin_bar.png"),
                     max_w_in=7.5, max_h_in=4.9, top_in=1.3, center_x_in=6.667)
    add_bullets(slide, 0.6, 6.35, 11.6, 0.8, [
        ("|phi|>=90 문제와는 무관(MAE 6.54mm vs 6.17mm로 거의 동일) — s값 자체가 클수록(팁 쪽) "
         "FEA 신호가 약해지는 옛 패턴과 일치. 다음 후보: 60-100mm 구간 실측 FEA 추가 확보", False),
    ], size_pt=14)

    # ---- Slide 6: L_M=0 문제 (부분 성공) ----
    slide = new_content_slide(prs, "5. L_M=0 문제 — 하이브리드 구조로 시도, 부분적 성공", 6)
    add_picture_fit(slide, os.path.join(CS_DIR, "lm_zero_hybrid_bar.png"),
                     max_w_in=7.5, max_h_in=4.9, top_in=1.3, center_x_in=6.667)
    add_bullets(slide, 0.6, 6.35, 11.6, 0.8, [
        ("\"0인지 아닌지\" 분류는 쉬운데(정확도 96%) \"정확히 몇 mm\" 연속회귀만 유독 실패하는 걸 "
         "확인 -> 분류+회귀 2단계 구조로 전환, 방향은 항상 개선되지만 개선 폭은 시드마다 다름", False),
    ], size_pt=14)

    # ---- Slide 7: 다음 계획 ----
    slide = new_content_slide(prs, "6. 다음 계획", 7)
    add_bullets(slide, 0.7, 1.6, 11.0, 4.5, [
        ("HIGH_PHI_WEIGHT 레버는 접고, 검증된 방법인 실측 FEA 데이터 확보로 방향 전환", True),
        ("phi=90~150° 구간 실측 FEA 추가 스윕 — Fx_board 고각도 문제의 근본 원인(그 구간 데이터 "
         "자체가 적음)을 직접 해결", False),
        ("s=60~100mm(팁 근처) 구간 실측 FEA 추가 확보 — 같은 종류의 스윕 작업이라 위와 함께 진행 가능", False),
        ("L_M=0 하이브리드는 여러 시드로 재확인해서 개선 폭의 안정성 확보(급하지 않음)", False),
    ], size_pt=17)

    prs.save(OUT_PATH)
    print("saved", OUT_PATH)


if __name__ == "__main__":
    build()
