"""8월 2째주 랩미팅 pptx 조립.

스타일(색상/폰트크기/여백)은 기존 '랩미팅_8월1째주 (2).pptx'에서 그대로 추출한 값 사용.
내용은 사진 위주, 텍스트는 슬라이드 제목 + 페이지 번호뿐.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

GREEN = RGBColor(0x27, 0xAE, 0x60)
TITLE_DARK = RGBColor(0x22, 0x22, 0x22)
NAVY = RGBColor(0x1F, 0x2A, 0x44)
GRAY = RGBColor(0x6B, 0x72, 0x80)

EMU_PER_IN = 914400
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

HERE = os.path.dirname(__file__)
FEA_DIR = os.path.join(HERE, "..", "..", "data", "contact_scenarios", "fea")
OUT_PATH = r"C:\Users\gustj\Desktop\김현서\랩미팅 발표 자료\랩미팅_8월2째주.pptx"


def in_(v):
    return Emu(int(v * EMU_PER_IN))


def set_run(run, size_pt, color, bold=False, font=None):
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    if font:
        run.font.name = font


def add_bottom_line(slide):
    line = slide.shapes.add_connector(1, in_(0.4), in_(7.1484), in_(12.9), in_(7.1484))
    line.line.color.rgb = GREEN
    line.line.width = Emu(19050)


def add_page_number(slide, n):
    tb = slide.shapes.add_textbox(in_(11.85), in_(7.05), in_(0.9), in_(0.35))
    tf = tb.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    run = tf.paragraphs[0].add_run()
    run.text = str(n)
    set_run(run, 11, GRAY)


def add_title_bar(slide, text):
    rect = slide.shapes.add_shape(1, in_(0.45), in_(0.35), in_(0.09), in_(0.55))
    rect.fill.solid()
    rect.fill.fore_color.rgb = GREEN
    rect.line.fill.background()

    tb = slide.shapes.add_textbox(in_(0.65), in_(0.28), in_(11.5), in_(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_run(run, 26, TITLE_DARK, bold=True)


def new_content_slide(prs, title, page_num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, title)
    add_bottom_line(slide)
    add_page_number(slide, page_num)
    return slide


def add_picture_fit(slide, path, max_w_in, max_h_in, top_in, center_x_in=6.667):
    from PIL import Image
    w_px, h_px = Image.open(path).size
    aspect = w_px / h_px
    w_in, h_in = max_w_in, max_w_in / aspect
    if h_in > max_h_in:
        h_in = max_h_in
        w_in = max_h_in * aspect
    left = center_x_in - w_in / 2
    slide.shapes.add_picture(path, in_(left), in_(top_in), width=in_(w_in), height=in_(h_in))


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ---- Title slide ----
    slide = prs.slides.add_slide(blank)
    rect = slide.shapes.add_shape(1, in_(0.5), in_(0.45), in_(0.35), in_(0.7))
    rect.fill.solid()
    rect.fill.fore_color.rgb = GREEN
    rect.line.fill.background()

    tb = slide.shapes.add_textbox(in_(0.5), in_(2.6), in_(9), in_(0.85))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "8월 2째주 랩미팅"
    set_run(run, 44, GREEN, bold=True)

    tb = slide.shapes.add_textbox(in_(0.5), in_(3.5), in_(10), in_(0.6))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "구간 인지 모델 — 실전 조건(싱글프로브) 성능 & 접촉각(β) 전환점 원인 규명"
    set_run(run, 18, TITLE_DARK)

    tb = slide.shapes.add_textbox(in_(0.5), in_(4.6), in_(10), in_(0.5))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "20242260 김현서"
    set_run(run, 20, NAVY, bold=True)

    tb = slide.shapes.add_textbox(in_(0.5), in_(5.1), in_(10), in_(0.5))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "전기전자공학전공"
    set_run(run, 14, GRAY)

    add_bottom_line(slide)

    # ---- Content slides ----
    pages = [
        ("1. 지난주 리캡 & 이번 주 질문", "review0813_recap_questions.png"),
        ("2. β(접촉각)란? — 그림자 비유로 이해하기", "review0813_beta_explainer.png"),
        ("3. β 전환점 정밀 규명 (실측 데이터)", "review0813_beta_transition.png"),
        ("4. 다른 형상에서도 재현되는가", "review0813_beta_generalization.png"),
        ("5. 프로브 개수 실험 히스토리 — 왜 싱글프로브를 다시 봤나", "review0813_probe_history.png"),
        ("6. 싱글프로브란 — 무엇이 바뀐 조건인가", "review0813_singleprobe_setup.png"),
        ("7. 싱글프로브 결과 상세", "review0813_singleprobe_breakdown.png"),
        ("8. 자연스러운 움직임 재활용이란?", "review0813_naturalmotion_concept.png"),
        ("9. 자연스러운 움직임 재활용 — 결과 비교", "review0813_naturalmotion_comparison.png"),
        ("10. 자연스러운 움직임 재활용 — 결과 상세", "review0813_naturalmotion_breakdown.png"),
        ("11. 견고성 체크 — 모델 구조 때문 아님", "review0813_ablation_robustness.png"),
        ("12. 이번 주 접근이 타당한 이유", "review0813_rationale_roadmap.png"),
    ]
    for i, (title, fname) in enumerate(pages, start=2):
        slide = new_content_slide(prs, title, i)
        add_picture_fit(slide, os.path.join(FEA_DIR, fname), max_w_in=11.6, max_h_in=5.6, top_in=1.35)

    prs.save(OUT_PATH)
    print("saved", OUT_PATH)


if __name__ == "__main__":
    build()
