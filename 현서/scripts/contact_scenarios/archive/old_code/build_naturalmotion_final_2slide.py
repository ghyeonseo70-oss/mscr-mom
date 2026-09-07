"""3-probe(자연스러운 움직임 재활용) 15만개 최종 결과 - 2장짜리 요약 pptx.
스타일은 build_lab_meeting_0813.py와 동일(같은 헬퍼 재사용).
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

GREEN = RGBColor(0x27, 0xAE, 0x60)
TITLE_DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x6B, 0x72, 0x80)

EMU_PER_IN = 914400
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

HERE = os.path.dirname(__file__)
FEA_DIR = os.path.join(HERE, "fea", "..", "..", "..", "data", "contact_scenarios", "fea")
OUT_PATH = r"C:\Users\gustj\Desktop\김현서\랩미팅 발표 자료\3프로브_최종결과_15만개.pptx"


def in_(v):
    return Emu(int(v * EMU_PER_IN))


def set_run(run, size_pt, color, bold=False):
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold


def add_bottom_line(slide):
    line = slide.shapes.add_connector(1, in_(0.4), in_(7.1484), in_(12.9), in_(7.1484))
    line.line.color.rgb = GREEN
    line.line.width = Emu(19050)


def add_page_number(slide, n):
    tb = slide.shapes.add_textbox(in_(11.85), in_(7.05), in_(0.9), in_(0.35))
    tf = tb.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    run = tf.paragraphs[0].add_run()
    run.text = str(n)
    set_run(run, 11, GRAY)


def add_title_bar(slide, text):
    rect = slide.shapes.add_shape(1, in_(0.45), in_(0.35), in_(0.09), in_(0.55))
    rect.fill.solid()
    rect.fill.fore_color.rgb = GREEN
    rect.line.fill.background()

    tb = slide.shapes.add_textbox(in_(0.65), in_(0.28), in_(11.5), in_(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_run(run, 26, TITLE_DARK, bold=True)


def new_content_slide(prs, title, page_num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, title)
    add_bottom_line(slide)
    add_page_number(slide, page_num)
    return slide


def add_picture_fit(slide, path, max_w_in, max_h_in, top_in, center_x_in=6.667):
    from PIL import Image
    w_px, h_px = Image.open(path).size
    aspect = w_px / h_px
    w_in, h_in = max_w_in, max_w_in / aspect
    if h_in > max_h_in:
        h_in = max_h_in
        w_in = max_h_in * aspect
    left = center_x_in - w_in / 2
    slide.shapes.add_picture(path, in_(left), in_(top_in), width=in_(w_in), height=in_(h_in))


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    pages = [
        ("3-probe(자연스러운 움직임 재활용) 15만개 최종 결과 — 11-probe 수준에 근접",
         "handoff_naturalmotion_final_comparison.png"),
        ("최종 결과 스코어카드 & 규모 확장 효과", "handoff_naturalmotion_final_scorecard.png"),
    ]
    for i, (title, fname) in enumerate(pages, start=1):
        slide = new_content_slide(prs, title, i)
        add_picture_fit(slide, os.path.join(FEA_DIR, fname), max_w_in=11.6, max_h_in=5.6, top_in=1.35)

    prs.save(OUT_PATH)
    print("saved", OUT_PATH)


if __name__ == "__main__":
    build()
