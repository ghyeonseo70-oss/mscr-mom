"""8월 3째주 랩미팅 pptx 조립 - 큰 오류 2개(서로게이트 환각학습, 2구간/3구간 착오) 중심으로
"지금까지 한 결과"만 정리. 스타일은 기존 '랩미팅_8월1째주 (2).pptx'에서 추출한 값 그대로 사용.
"""
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

GREEN = RGBColor(0x27, 0xAE, 0x60)
TITLE_DARK = RGBColor(0x22, 0x22, 0x22)
NAVY = RGBColor(0x1F, 0x2A, 0x44)
GRAY = RGBColor(0x6B, 0x72, 0x80)
BOX_FILL = RGBColor(0xEA, 0xF3, 0xEC)
RED = RGBColor(0xC0, 0x39, 0x2B)

EMU_PER_IN = 914400
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)

HERE = os.path.dirname(__file__)
FEA_DIR = os.path.join(HERE, "..", "..", "data", "contact_scenarios", "fea")
FORCE_MODEL_DIR = os.path.join(HERE, "..", "..", "..", "data", "force_model")
OUT_PATH = r"C:\Users\gustj\Desktop\김현서\랩미팅 발표 자료\랩미팅_8월3째주.pptx"


def in_(v):
    return Emu(int(v * EMU_PER_IN))


def set_run(run, size_pt, color, bold=False, font=None, align=None):
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.bold = bold
    if font:
        run.font.name = font


def add_bottom_line(slide):
    line = slide.shapes.add_connector(1, in_(0.4), in_(7.1484), in_(12.9), in_(7.1484))
    line.line.color.rgb = GREEN
    line.line.width = Emu(19050)


def add_page_number(slide, n):
    tb = slide.shapes.add_textbox(in_(11.85), in_(7.05), in_(0.9), in_(0.35))
    tf = tb.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    run = tf.paragraphs[0].add_run()
    run.text = str(n)
    set_run(run, 11, GRAY)


def add_title_bar(slide, text):
    rect = slide.shapes.add_shape(1, in_(0.45), in_(0.35), in_(0.09), in_(0.55))
    rect.fill.solid()
    rect.fill.fore_color.rgb = GREEN
    rect.line.fill.background()

    tb = slide.shapes.add_textbox(in_(0.65), in_(0.28), in_(11.5), in_(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    set_run(run, 26, TITLE_DARK, bold=True)


def new_content_slide(prs, title, page_num):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title_bar(slide, title)
    add_bottom_line(slide)
    add_page_number(slide, page_num)
    return slide


def add_picture_fit(slide, path, max_w_in, max_h_in, top_in, center_x_in=6.667):
    from PIL import Image
    w_px, h_px = Image.open(path).size
    aspect = w_px / h_px
    w_in, h_in = max_w_in, max_w_in / aspect
    if h_in > max_h_in:
        h_in = max_h_in
        w_in = max_h_in * aspect
    left = center_x_in - w_in / 2
    slide.shapes.add_picture(path, in_(left), in_(top_in), width=in_(w_in), height=in_(h_in))


def add_box(slide, left_in, top_in, w_in, h_in, text, fill=BOX_FILL, text_color=TITLE_DARK,
            size_pt=15, bold=False, line_color=GREEN):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, in_(left_in), in_(top_in), in_(w_in), in_(h_in))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = line_color
    box.line.width = Emu(19050)
    tf = box.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = line
        set_run(run, size_pt, text_color, bold=bold)
    return box


def add_arrow(slide, left_in, top_in, w_in, h_in):
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, in_(left_in), in_(top_in), in_(w_in), in_(h_in))
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = GREEN
    arrow.line.fill.background()
    return arrow


def add_bullets(slide, left_in, top_in, w_in, h_in, lines, size_pt=17):
    tb = slide.shapes.add_textbox(in_(left_in), in_(top_in), in_(w_in), in_(h_in))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (line, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = "• " + line
        set_run(run, size_pt, TITLE_DARK, bold=bold)
        p.space_after = Pt(10)


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # ---- Title slide ----
    slide = prs.slides.add_slide(blank)
    rect = slide.shapes.add_shape(1, in_(0.5), in_(0.45), in_(0.35), in_(0.7))
    rect.fill.solid()
    rect.fill.fore_color.rgb = GREEN
    rect.line.fill.background()

    tb = slide.shapes.add_textbox(in_(0.5), in_(2.6), in_(11), in_(0.85))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "8월 3째주 랩미팅"
    set_run(run, 44, GREEN, bold=True)

    tb = slide.shapes.add_textbox(in_(0.5), in_(3.5), in_(11), in_(0.6))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "힘 추정이 무너진 이유와 회복 — 실측 검증에서 드러난 두 가지 큰 오류"
    set_run(run, 18, TITLE_DARK)

    tb = slide.shapes.add_textbox(in_(0.5), in_(4.6), in_(10), in_(0.5))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "20242260 김현서"
    set_run(run, 20, NAVY, bold=True)

    tb = slide.shapes.add_textbox(in_(0.5), in_(5.1), in_(10), in_(0.5))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "전기전자공학전공"
    set_run(run, 14, GRAY)

    add_bottom_line(slide)

    # ---- Slide 2: 문제 발견 ----
    slide = new_content_slide(prs, "1. 문제 발견 — 합성 검증에서는 좋아 보였는데", 2)
    add_picture_fit(slide, os.path.join(FEA_DIR, "lab0821_r2_collapse.png"),
                     max_w_in=9.0, max_h_in=5.6, top_in=1.35)

    # ---- Slide 3: 원인 - 서로게이트 환각 학습 (native diagram) ----
    slide = new_content_slide(prs, "2. 원인 — \"서로게이트 환각 학습\"", 3)
    box_w, box_h, gap = 3.3, 1.7, 0.55
    arrow_w = gap - 0.1
    top = 1.7
    x0 = 0.6
    add_box(slide, x0, top, box_w, box_h,
            "실측 변위 예측이 약함\n(tip_ux/uy R²=0.52~0.68)",
            fill=RGBColor(0xFC, 0xEF, 0xE3))
    add_arrow(slide, x0 + box_w, top + box_h/2 - 0.2, arrow_w, 0.4)
    add_box(slide, x0 + box_w + gap, top, box_w, box_h,
            "CNN이 진짜 물리 대신\n대체모델의 (부정확한)\n변위→힘 매핑을 그대로 배움",
            fill=BOX_FILL)
    add_arrow(slide, x0 + 2*box_w + gap, top + box_h/2 - 0.2, arrow_w, 0.4)
    add_box(slide, x0 + 2*(box_w + gap), top, box_w, box_h,
            "실측 힘 R²=0.337로 붕괴\n(원래 대체모델 단독으로는\nR²=0.94로 잘 배우던 타겟인데도)",
            fill=RGBColor(0xFC, 0xE4, 0xE1), text_color=RED)
    add_bullets(slide, 0.6, top + box_h + 0.6, 10.8, 2.0, [
        ("확인 방법: 대체모델 스타일로 만든 입력을 주면 대체모델 자신의 예측과 R²=0.896으로 "
         "거의 그대로 일치 → CNN이 물리가 아니라 대체모델의 매핑 자체를 외운 것", False),
        ("데이터 없이 고치려는 시도(피처 엔지니어링, 물리모델로 변위 대체) 2개 모두 효과 없었음 "
         "→ 실측 FEA를 늘려서 변위 예측을 개선하는 것만이 남은 해법이었음", False),
    ], size_pt=15)

    # ---- Slide 4: 해결 ----
    slide = new_content_slide(prs, "3. 해결 — 실측 FEA 데이터 확대 (L_M 조밀화)", 4)
    add_picture_fit(slide, os.path.join(FEA_DIR, "lab0821_r2_recovery.png"),
                     max_w_in=9.0, max_h_in=5.6, top_in=1.35)

    # ---- Slide 5: 또 다른 큰 오류 - 2구간/3구간 착오 ----
    slide = new_content_slide(prs, "4. 또 다른 큰 오류 — 2구간 vs 3구간 구조 착오", 5)
    add_picture_fit(slide, os.path.join(FORCE_MODEL_DIR, "fig3_full_8panel_reproduction.png"),
                     max_w_in=11.6, max_h_in=4.5, top_in=1.3)
    add_bullets(slide, 0.6, 5.95, 11.0, 1.1, [
        ("MOM(8mm) 구간을 강체 직선으로 보는 K_RIGID를 실수로 빼고 2구간 모델로 바꿨다가, "
         "논문 원문(식 9 부근: \"자석 길이를 직선으로 취급\")을 다시 확인하고 3구간으로 원복함 "
         "— 위 그림처럼 논문 Fig.3(a)-(h)와 다시 잘 맞음", False),
    ], size_pt=14)

    # ---- Slide 6: 다음 계획 ----
    slide = new_content_slide(prs, "5. 다음 계획", 6)
    add_bullets(slide, 0.7, 1.6, 11.0, 4.5, [
        ("힘 R²는 회복됐지만(0.65~0.73) 아직 0.9대는 아님 — phi=90~150 구간 FEA 성공률이 "
         "여전히 낮아서(30~40%대) 그 구간 데이터를 더 늘리는 게 다음 후보", True),
        ("전략 1: 기존 s격자(10,20,...,100mm) 사이 오프셋 격자로 새 샘플 추가 시도 (378케이스)", False),
        ("전략 2: STABILIZE 감쇠계수를 자동값 대신 명시적으로 튜닝해서 성공률 자체를 올릴 수 "
         "있는지 소규모 파일럿(60케이스)으로 확인", False),
        ("현재 다른 컴퓨터에서 두 스윕 실행 대기 중 — 끝나면 재학습해서 힘 R²가 더 오르는지 확인", False),
    ], size_pt=17)

    prs.save(OUT_PATH)
    print("saved", OUT_PATH)


if __name__ == "__main__":
    build()
